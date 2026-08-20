"""
Generador de Datos para Power BI - Seguimiento Nevasa
"""
import os, sys
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config.rutas import ARCHIVO_FUENTE, ARCHIVO_SALIDA_POWERBI, HH_DIARIAS

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    print("Error: openpyxl no esta instalado. Ejecuta: pip install openpyxl")
    sys.exit(1)


def networkdays(fecha_inicio, fecha_fin):
    if not fecha_inicio or not fecha_fin:
        return None
    if not isinstance(fecha_inicio, datetime) or not isinstance(fecha_fin, datetime):
        return None
    inicio = fecha_inicio.date() if isinstance(fecha_inicio, datetime) else fecha_inicio
    fin = fecha_fin.date() if isinstance(fecha_fin, datetime) else fecha_fin
    if inicio > fin:
        return 0
    dias = 0
    actual = inicio
    while actual <= fin:
        if actual.weekday() < 5:
            dias += 1
        actual += timedelta(days=1)
    return dias


def leer_fuente():
    print(f"Leyendo fuente: {ARCHIVO_FUENTE}")
    wb = openpyxl.load_workbook(ARCHIVO_FUENTE, data_only=True)
    ws = wb["Seguimiento Nevasa"]
    filas = []
    for r in range(5, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(x is None for x in v):
            continue

        idx = v[0]
        req = v[1] or ""
        det = v[2] or ""
        area = v[3] or ""
        sol = v[4] or ""
        cat = v[5] or ""
        pri_raw = v[6]
        comp = v[7] or ""
        asig = v[8] or ""
        etapa = v[9] or "Pendiente"
        apli = v[10] or ""
        f_rec = v[11]
        ei = v[12]
        eq = v[13]
        ef = v[14]
        te = v[15]
        de_d = v[16]
        ri = v[16+1]
        frq = v[18]
        pap = v[19] or ""
        fp = v[20]
        h_real = v[22]
        dr_d = v[24]
        hist = v[25] or ""
        plan = v[26] or ""
        periodo = v[27]

        if pri_raw is None or pri_raw == "" or pri_raw == "Sin Prioridad":
            pri = "Sin Prioridad"
        elif pri_raw == 0 or pri_raw == "0":
            pri = "Maxima"
        elif pri_raw == 1 or pri_raw == "1":
            pri = "Alta"
        else:
            pri = str(pri_raw)

        de = None
        if de_d and isinstance(de_d, (int, float)):
            de = int(de_d)
        elif ei and ef and isinstance(ei, datetime) and isinstance(ef, datetime):
            de = networkdays(ei, ef)

        dr = None
        if dr_d and isinstance(dr_d, (int, float)):
            dr = int(dr_d)
        elif ri and fp and isinstance(ri, datetime) and isinstance(fp, datetime):
            dr = networkdays(ri, fp)

        he = de * HH_DIARIAS if de else None
        hr = dr * HH_DIARIAS if dr else None

        cola_e = networkdays(f_rec, ei) if f_rec and ei and isinstance(f_rec, datetime) and isinstance(ei, datetime) else None
        cola_r = networkdays(f_rec, ri) if f_rec and ri and isinstance(f_rec, datetime) and isinstance(ri, datetime) else None
        desa_e = networkdays(ei, ef) if ei and ef and isinstance(ei, datetime) and isinstance(ef, datetime) else None
        desa_r = networkdays(ri, fp) if ri and fp and isinstance(ri, datetime) and isinstance(fp, datetime) else None
        qa_r = networkdays(frq, fp) if frq and fp and isinstance(frq, datetime) and isinstance(fp, datetime) else None

        var_d = (dr - de) if de is not None and dr is not None else None
        var_p = round((var_d / de) * 100, 1) if de and de > 0 and var_d is not None else None
        a_tiempo = "Si" if de is not None and dr is not None and dr <= de else ("No" if de is not None and dr is not None else None)
        ciclo_r = networkdays(f_rec, fp) if f_rec and fp and isinstance(f_rec, datetime) and isinstance(fp, datetime) else None

        filas.append({
            "Indice": idx,
            "Requerimiento": req,
            "Detalle": det,
            "Area": area,
            "Solicitante": sol,
            "Categoria": cat,
            "Prioridad": pri,
            "Complejidad": comp,
            "Asignado_a": asig,
            "Etapa": etapa,
            "Aplicativo": apli,
            "Fecha_Recepcion": f_rec,
            "Estimado_Inicio": ei,
            "Estimado_QA": eq,
            "Estimado_Fin": ef,
            "Real_Inicio": ri,
            "Fecha_Real_QA": frq,
            "Fecha_Produccion": fp,
            "Periodo": periodo,
            "Aprobado_PaP": pap,
            "Historia": hist,
            "Plan_Prueba": plan,
            "Duracion_Estimada": de,
            "Duracion_Real": dr,
            "Horas_Estimadas": he,
            "Horas_Reales": hr,
            "Cola_Estimada": cola_e,
            "Cola_Real": cola_r,
            "Desarrollo_Estimado": desa_e,
            "Desarrollo_Real": desa_r,
            "QA_Real": qa_r,
            "Varianza_Dias": var_d,
            "Varianza_Pct": var_p,
            "A_Tiempo": a_tiempo,
            "Ciclo_Total": ciclo_r,
        })
    wb.close()
    print(f"  Filas leidas: {len(filas)}")
    return filas


def escribir_excel(filas):
    print(f"\nGenerando: {ARCHIVO_SALIDA_POWERBI}")
    wb = Workbook()

    ws = wb.active
    ws.title = "Datos"

    headers = [
        "Indice", "Requerimiento", "Detalle", "Area", "Solicitante",
        "Categoria", "Prioridad", "Complejidad", "Asignado_a", "Etapa",
        "Aplicativo", "Fecha_Recepcion", "Estimado_Inicio", "Estimado_QA",
        "Estimado_Fin", "Real_Inicio", "Fecha_Real_QA", "Fecha_Produccion",
        "Periodo", "Aprobado_PaP", "Historia", "Plan_Prueba",
        "Duracion_Estimada", "Duracion_Real", "Horas_Estimadas", "Horas_Reales",
        "Cola_Estimada", "Cola_Real", "Desarrollo_Estimado", "Desarrollo_Real",
        "QA_Real", "Varianza_Dias", "Varianza_Pct", "A_Tiempo", "Ciclo_Total"
    ]

    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    thin = Side(style="thin")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    date_cols = {12, 13, 14, 15, 16, 17, 18, 19}
    num_cols = {23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 35}

    for r_idx, f in enumerate(filas, 2):
        for c, h in enumerate(headers, 1):
            val = f.get(h)
            if h in date_cols and isinstance(val, datetime):
                val = val.date()
            cell = ws.cell(r_idx, c, val)
            cell.border = border
            if h in date_cols:
                cell.number_format = "DD/MM/YYYY"
                cell.alignment = Alignment(horizontal="center")
            elif h in num_cols:
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center")
            elif h == 33:
                cell.number_format = "0.0"
                cell.alignment = Alignment(horizontal="center")

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 16

    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["U"].width = 50

    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{len(filas) + 1}"
    ws.freeze_panes = "A2"

    ws_dim = wb.create_sheet("Dim_Area")
    areas = sorted(set(f["Area"] for f in filas if f["Area"]))
    ws_dim.append(["Area"])
    for a in areas:
        ws_dim.append([a])

    ws_dim2 = wb.create_sheet("Dim_Etapa")
    etapas = ["Sin Formalizar", "Pendiente", "Analisis", "Desarrollo",
              "Certificacion Creasys", "Certificacion Nevasa", "Produccion",
              "Stand By", "Cancelado"]
    ws_dim2.append(["Etapa", "Orden"])
    for i, e in enumerate(etapas):
        ws_dim2.append([e, i])

    ws_dim3 = wb.create_sheet("Dim_Prioridad")
    pris = ["Maxima", "Alta", "Media/Alta", "Media", "Baja", "Sin Prioridad"]
    ws_dim3.append(["Prioridad", "Orden"])
    for i, p in enumerate(pris):
        ws_dim3.append([p, i])

    print(f"  Hojas base generadas: Datos, Dim_Area, Dim_Etapa, Dim_Prioridad")
    return wb


def agregar_hoja_riesgos_powerbi(wb, archivo_fuente):
    """Agrega tabla de riesgos para Power BI."""
    try:
        import openpyxl
        wb_src = openpyxl.load_workbook(archivo_fuente, read_only=True)
        if "Riesgos" not in wb_src.sheetnames:
            wb_src.close()
            return False

        ws_src = wb_src["Riesgos"]
        ws_dst = wb.create_sheet("Riesgos")

        headers = ["ID", "Descripcion", "Probabilidad", "Impacto", "Estado",
                    "Mitigacion", "Responsable", "Fecha_Identificacion", "Fecha_Cierre", "Notas"]
        ws_dst.append(headers)

        for r in range(2, ws_src.max_row + 1):
            valores = [ws_src.cell(r, c).value for c in range(1, 11)]
            if all(v is None for v in valores):
                continue
            ws_dst.append(valores)

        wb_src.close()
        return True
    except Exception:
        return False


def agregar_hoja_lessons_powerbi(wb, archivo_fuente):
    """Agrega tabla de lessons learned para Power BI."""
    try:
        import openpyxl
        wb_src = openpyxl.load_workbook(archivo_fuente, read_only=True)
        if "Lessons Learned" not in wb_src.sheetnames:
            wb_src.close()
            return False

        ws_src = wb_src["Lessons Learned"]
        ws_dst = wb.create_sheet("Lessons")

        headers = ["ID", "Fecha", "Categoria", "Descripcion", "Impacto",
                    "Leccion", "Accion_Correctiva", "Aplicable_a", "Registrado_por"]
        ws_dst.append(headers)

        for r in range(2, ws_src.max_row + 1):
            valores = [ws_src.cell(r, c).value for c in range(1, 10)]
            if all(v is None for v in valores):
                continue
            ws_dst.append(valores)

        wb_src.close()
        return True
    except Exception:
        return False


def agregar_hoja_recursos_powerbi(wb, archivo_fuente):
    """Agrega tabla de recursos para Power BI."""
    try:
        import openpyxl
        wb_src = openpyxl.load_workbook(archivo_fuente, read_only=True)
        if "Recursos" not in wb_src.sheetnames:
            wb_src.close()
            return False

        ws_src = wb_src["Recursos"]
        ws_dst = wb.create_sheet("Recursos")

        headers = ["Persona", "Horas_Estimadas", "Horas_Reales", "Capacidad_Semanal",
                    "Fase_Actual", "Asignado_a"]
        ws_dst.append(headers)

        for r in range(2, ws_src.max_row + 1):
            valores = [ws_src.cell(r, c).value for c in range(1, 7)]
            if all(v is None for v in valores):
                continue
            ws_dst.append(valores)

        wb_src.close()
        return True
    except Exception:
        return False


def main():
    filas = leer_fuente()
    if not filas:
        print("No se encontraron datos.")
        return
    wb = escribir_excel(filas)

    # Agregar tablas adicionales si existen
    from config.rutas import ARCHIVO_FUENTE
    agregar_hoja_riesgos_powerbi(wb, ARCHIVO_FUENTE)
    agregar_hoja_lessons_powerbi(wb, ARCHIVO_FUENTE)
    agregar_hoja_recursos_powerbi(wb, ARCHIVO_FUENTE)

    wb.save(ARCHIVO_SALIDA_POWERBI)

    print("\n=== RESUMEN ===")
    total = len(filas)
    completados = sum(1 for f in filas if f["Etapa"] == "Produccion")
    en_prog = sum(1 for f in filas if f["Etapa"] not in ("Produccion", "Cancelado", "Pendiente"))
    pend = sum(1 for f in filas if f["Etapa"] == "Pendiente")
    print(f"  Total: {total}")
    print(f"  Completados: {completados}")
    print(f"  En Progreso: {en_prog}")
    print(f"  Pendientes: {pend}")
    print(f"\n  Para importar a Power BI:")
    print(f"  1. Abrir Power BI Desktop")
    print(f"  2. Obtener datos > Excel")
    print(f"  3. Seleccionar: {ARCHIVO_SALIDA_POWERBI}")
    print(f"  4. Cargar tablas: Datos, Riesgos, Lessons, Recursos")


if __name__ == "__main__":
    main()
