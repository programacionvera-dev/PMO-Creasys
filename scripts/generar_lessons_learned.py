"""
Generador de Lessons Learned
Lee la hoja 'Lessons Learned' del Excel fuente y genera reporte de lecciones.
"""
import sys
import os
from datetime import datetime
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config.rutas import (
    ARCHIVO_FUENTE, ARCHIVO_SALIDA_REPORTE, MAPEO_LESSONS,
    CATEGORIAS_LESSON, BRANDING
)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Colores
FILL_HEADER = PatternFill(start_color="FF333F48", end_color="FF333F48", fill_type="solid")
FILL_TECNICO = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
FILL_PROCESO = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
FILL_COMUNICACION = PatternFill(start_color="FFE8E8E8", end_color="FFE8E8E8", fill_type="solid")
FILL_RECURSO = PatternFill(start_color="FFF8DC", end_color="FFF8DC", fill_type="solid")
FILL_CLIENTE = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")

FONT_HEADER = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
FONT_BODY = Font(name="Calibri", size=9)
FONT_TITULO = Font(name="Calibri", size=16, bold=True, color="FF333F48")
FONT_SECCION = Font(name="Calibri", size=11, bold=True)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="FFB0B0B0"),
    right=Side(style="thin", color="FFB0B0B0"),
    top=Side(style="thin", color="FFB0B0B0"),
    bottom=Side(style="thin", color="FFB0B0B0"),
)

CATEGORIA_FILL_MAP = {
    "Técnico": FILL_TECNICO,
    "Proceso": FILL_PROCESO,
    "Comunicación": FILL_COMUNICACION,
    "Recurso": FILL_RECURSO,
    "Cliente": FILL_CLIENTE,
}


def leer_lessons():
    """Lee la hoja 'Lessons Learned' del Excel fuente."""
    try:
        wb = openpyxl.load_workbook(ARCHIVO_FUENTE, read_only=True)
    except FileNotFoundError:
        print(f"Archivo no encontrado: {ARCHIVO_FUENTE}")
        return []
    except KeyError:
        print("Hoja 'Lessons Learned' no encontrada en el Excel.")
        return []

    if "Lessons Learned" not in wb.sheetnames:
        print("Hoja 'Lessons Learned' no encontrada en el Excel.")
        wb.close()
        return []

    ws = wb["Lessons Learned"]
    m = MAPEO_LESSONS
    lessons = []

    for r in range(2, ws.max_row + 1):
        valores = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in valores):
            continue

        lesson = {
            "id": valores[m["id"]] if len(valores) > m["id"] else None,
            "fecha": valores[m["fecha"]] if len(valores) > m["fecha"] else None,
            "categoria": valores[m["categoria"]] if len(valores) > m["categoria"] else "",
            "descripcion": valores[m["descripcion"]] if len(valores) > m["descripcion"] else "",
            "impacto": valores[m["impacto"]] if len(valores) > m["impacto"] else "",
            "leccion": valores[m["leccion"]] if len(valores) > m["leccion"] else "",
            "accion_correctiva": valores[m["accion_correctiva"]] if len(valores) > m["accion_correctiva"] else "",
            "aplicable_a": valores[m["aplicable_a"]] if len(valores) > m["aplicable_a"] else "",
            "registrado_por": valores[m["registrado_por"]] if len(valores) > m["registrado_por"] else "",
        }
        lessons.append(lesson)

    wb.close()
    return lessons


def calcular_metricas(lessons):
    """Calcula métricas de lessons learned."""
    total = len(lessons)

    por_categoria = defaultdict(int)
    con_accion = 0
    for l in lessons:
        cat = l["categoria"] if l["categoria"] else "Sin categoría"
        por_categoria[cat] += 1
        if l["accion_correctiva"] and str(l["accion_correctiva"]).strip():
            con_accion += 1

    return {
        "total": total,
        "por_categoria": dict(por_categoria),
        "con_accion_correctiva": con_accion,
        "porcentaje_accion": round((con_accion / total * 100)) if total > 0 else 0,
    }


def _aplicar_estilo(celda, font=None, fill=None, alignment=None, border=None):
    if font:
        celda.font = font
    if fill:
        celda.fill = fill
    if alignment:
        celda.alignment = alignment
    if border:
        celda.border = border


def _merge_y_escribir(ws, r1, c1, r2, c2, valor, font, alignment=None, fill=None):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    celda = ws.cell(r1, c1, valor)
    _aplicar_estilo(celda, font=font, alignment=alignment or Alignment(horizontal="left"), fill=fill)


def agregar_hoja_lessons(wb, lessons, metricas, fecha_gen):
    """Agrega la hoja de lessons learned al workbook."""
    ws = wb.create_sheet("Lessons Learned")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 40
    ws.column_dimensions["G"].width = 40
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["I"].width = 15

    # Título
    _merge_y_escribir(ws, 1, 1, 1, 9, "LESSONS LEARNED", FONT_TITULO, ALIGN_CENTER)
    _merge_y_escribir(ws, 2, 1, 2, 9, f"Fecha: {fecha_gen}", Font(name="Calibri", size=10, color="FF97999B"), ALIGN_CENTER)

    # Resumen
    row = 4
    _merge_y_escribir(ws, row, 1, row, 9, "RESUMEN", FONT_SECCION)
    row += 1
    headers_resumen = ["Total Lecciones", "Con Acción Correctiva", "% Con Acción"]
    for c, h in enumerate(headers_resumen, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    valores = [metricas["total"], metricas["con_accion_correctiva"], metricas["porcentaje_accion"] / 100]
    for c, v in enumerate(valores, 1):
        celda = ws.cell(row, c, v)
        _aplicar_estilo(celda, font=FONT_BODY, border=THIN_BORDER, alignment=ALIGN_CENTER)
        if c == 3:
            celda.number_format = '0%'

    # Distribución por categoría
    row += 2
    _merge_y_escribir(ws, row, 1, row, 9, "DISTRIBUCIÓN POR CATEGORÍA", FONT_SECCION)
    row += 1
    for c, h in enumerate(["Categoría", "Cantidad", "% del Total"], 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    for cat in CATEGORIAS_LESSON:
        cant = metricas["por_categoria"].get(cat, 0)
        pct = cant / metricas["total"] if metricas["total"] > 0 else 0
        fill = CATEGORIA_FILL_MAP.get(cat)
        for c, v in enumerate([cat, cant, pct], 1):
            celda = ws.cell(row, c, v)
            _aplicar_estilo(celda, font=FONT_BODY, border=THIN_BORDER, fill=fill,
                           alignment=ALIGN_CENTER if c > 1 else None)
            if c == 3:
                celda.number_format = '0%'
        row += 1

    # Detalle de lecciones
    row += 1
    _merge_y_escribir(ws, row, 1, row, 9, "DETALLE DE LECCIONES", FONT_SECCION)
    row += 1
    headers = ["ID", "Fecha", "Categoría", "Descripción", "Impacto", "Lección",
               "Acción Correctiva", "Aplicable a", "Registrado por"]
    for c, h in enumerate(headers, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1

    for l in lessons:
        fill = CATEGORIA_FILL_MAP.get(l["categoria"])
        vals = [
            l["id"], l["fecha"], l["categoria"], l["descripcion"],
            l["impacto"], l["leccion"], l["accion_correctiva"],
            l["aplicable_a"], l["registrado_por"]
        ]
        for c, v in enumerate(vals, 1):
            celda = ws.cell(row, c, v)
            _aplicar_estilo(celda, font=FONT_BODY, border=THIN_BORDER, fill=fill,
                           alignment=ALIGN_CENTER if c in (1, 2, 3) else None)
        row += 1


def main():
    print("Leyendo lessons learned...")
    lessons = leer_lessons()

    if not lessons:
        print("No se encontraron lecciones. Asegúrese de que la hoja 'Lessons Learned' exista en el Excel.")
        return

    metricas = calcular_metricas(lessons)
    fecha_gen = datetime.now().strftime("%d/%m/%Y")

    print(f"  Total: {metricas['total']}")
    print(f"  Con acción correctiva: {metricas['con_accion_correctiva']}")
    print(f"  % Con acción: {metricas['porcentaje_accion']}%")

    try:
        wb = openpyxl.load_workbook(ARCHIVO_SALIDA_REPORTE)
    except FileNotFoundError:
        print(f"Archivo de reporte no encontrado: {ARCHIVO_SALIDA_REPORTE}")
        print("Ejecute primero generar_reporte_nevasa.py")
        return

    agregar_hoja_lessons(wb, lessons, metricas, fecha_gen)

    archivo_temp = ARCHIVO_SALIDA_REPORTE.replace(".xlsx", "_temp.xlsx")
    wb.save(archivo_temp)
    try:
        os.replace(archivo_temp, ARCHIVO_SALIDA_REPORTE)
    except PermissionError:
        print(f"\n  [!] No se pudo sobrescribir. Guardado como: {archivo_temp}")
        return
    print(f"\nLessons Learned agregado a: {ARCHIVO_SALIDA_REPORTE}")


if __name__ == "__main__":
    main()
