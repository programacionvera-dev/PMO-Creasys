"""
Generador de Análisis de Riesgos
Lee la hoja 'Riesgos' del Excel fuente y genera análisis de riesgos.
"""
import sys
import os
from datetime import datetime
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config.rutas import (
    ARCHIVO_FUENTE, ARCHIVO_SALIDA_REPORTE, MAPEO_RIESGOS,
    ESTADOS_RIESGO, UMBRALES_RIESGO, BRANDING
)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Colores
FILL_HEADER = PatternFill(start_color="FF333F48", end_color="FF333F48", fill_type="solid")
FILL_CRITICO = PatternFill(start_color="FFFF4444", end_color="FFFF4444", fill_type="solid")
FILL_ALTO = PatternFill(start_color="FFFF8C00", end_color="FFFF8C00", fill_type="solid")
FILL_MEDIO = PatternFill(start_color="FFFFD700", end_color="FFFFD700", fill_type="solid")
FILL_BAJO = PatternFill(start_color="FF90EE90", end_color="FF90EE90", fill_type="solid")
FILL_ABIERTO = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
FILL_MITIGADO = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
FILL_CERRADO = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")

FONT_HEADER = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
FONT_BODY = Font(name="Calibri", size=9)
FONT_TITULO = Font(name="Calibri", size=16, bold=True, color="FF333F48")
FONT_SECCION = Font(name="Calibri", size=11, bold=True)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="FFB0B0B0"),
    right=Side(style="thin", color="FFB0B0B0"),
    top=Side(style="thin", color="FFB0B0B0"),
    bottom=Side(style="thin", color="FFB0B0B0"),
)

PROB_MAP = {"Baja": 1, "Media": 2, "Alta": 3}
IMPACTO_MAP = {"Bajo": 1, "Medio": 2, "Alto": 3, "Crítico": 4}


def leer_riesgos():
    """Lee la hoja 'Riesgos' del Excel fuente."""
    try:
        wb = openpyxl.load_workbook(ARCHIVO_FUENTE, read_only=True)
    except FileNotFoundError:
        print(f"Archivo no encontrado: {ARCHIVO_FUENTE}")
        return []
    except KeyError:
        print("Hoja 'Riesgos' no encontrada en el Excel.")
        return []

    if "Riesgos" not in wb.sheetnames:
        print("Hoja 'Riesgos' no encontrada en el Excel.")
        wb.close()
        return []

    ws = wb["Riesgos"]
    m = MAPEO_RIESGOS
    riesgos = []

    for r in range(2, ws.max_row + 1):
        valores = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in valores):
            continue

        riesgo = {
            "id": valores[m["id"]] if len(valores) > m["id"] else None,
            "descripcion": valores[m["descripcion"]] if len(valores) > m["descripcion"] else "",
            "probabilidad": valores[m["probabilidad"]] if len(valores) > m["probabilidad"] else "",
            "impacto": valores[m["impacto"]] if len(valores) > m["impacto"] else "",
            "estado": valores[m["estado"]] if len(valores) > m["estado"] else "Abierto",
            "mitigacion": valores[m["mitigacion"]] if len(valores) > m["mitigacion"] else "",
            "responsable": valores[m["responsable"]] if len(valores) > m["responsable"] else "",
            "fecha_identificacion": valores[m["fecha_identificacion"]] if len(valores) > m["fecha_identificacion"] else None,
            "fecha_cierre": valores[m["fecha_cierre"]] if len(valores) > m["fecha_cierre"] else None,
            "notas": valores[m["notas"]] if len(valores) > m["notas"] else "",
        }
        riesgos.append(riesgo)

    wb.close()
    return riesgos


def calcular_score(probabilidad, impacto):
    """Calcula el score de riesgo (1-12)."""
    p = PROB_MAP.get(str(probabilidad), 1)
    i = IMPACTO_MAP.get(str(impacto), 1)
    return p * i


def clasificar_riesgo(score):
    """Clasifica el riesgo según el score."""
    if score >= UMBRALES_RIESGO["critico"]:
        return "Crítico", FILL_CRITICO
    elif score >= UMBRALES_RIESGO["alto"]:
        return "Alto", FILL_ALTO
    elif score >= UMBRALES_RIESGO["medio"]:
        return "Medio", FILL_MEDIO
    else:
        return "Bajo", FILL_BAJO


def calcular_metricas(riesgos):
    """Calcula métricas de riesgos."""
    total = len(riesgos)
    abiertos = sum(1 for r in riesgos if r["estado"] == "Abierto")
    mitigados = sum(1 for r in riesgos if r["estado"] == "Mitigado")
    cerrados = sum(1 for r in riesgos if r["estado"] == "Cerrado")

    scores = []
    for r in riesgos:
        score = calcular_score(r["probabilidad"], r["impacto"])
        clasificacion, _ = clasificar_riesgo(score)
        r["score"] = score
        r["clasificacion"] = clasificacion
        scores.append(score)

    criticos = sum(1 for r in riesgos if r["clasificacion"] == "Crítico")
    altos = sum(1 for r in riesgos if r["clasificacion"] == "Alto")

    return {
        "total": total,
        "abiertos": abiertos,
        "mitigados": mitigados,
        "cerrados": cerrados,
        "criticos": criticos,
        "altos": altos,
        "score_promedio": round(sum(scores) / len(scores), 1) if scores else 0,
    }


def _aplicar_estilo(celda, font=None, fill=None, alignment=None, border=None):
    if font:
        celda.font = font
    if fill:
        celda.fill = fill
    if alignment:
        celda.alignment = alignment
    if border:
        celda.border = border


def _merge_y_escribir(ws, r1, c1, r2, c2, valor, font, alignment=None, fill=None):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    celda = ws.cell(r1, c1, valor)
    _aplicar_estilo(celda, font=font, alignment=alignment or Alignment(horizontal="left"), fill=fill)


def agregar_hoja_riesgos(wb, riesgos, metricas, fecha_gen):
    """Agrega la hoja de análisis de riesgos al workbook."""
    ws = wb.create_sheet("Análisis de Riesgos")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 40
    ws.column_dimensions["I"].width = 15
    ws.column_dimensions["J"].width = 15
    ws.column_dimensions["K"].width = 30

    # Título
    _merge_y_escribir(ws, 1, 1, 1, 11, "ANÁLISIS DE RIESGOS", FONT_TITULO, ALIGN_CENTER)
    _merge_y_escribir(ws, 2, 1, 2, 11, f"Fecha: {fecha_gen}", Font(name="Calibri", size=10, color="FF97999B"), ALIGN_CENTER)

    # Resumen
    row = 4
    _merge_y_escribir(ws, row, 1, row, 11, "RESUMEN DE RIESGOS", FONT_SECCION)
    row += 1
    headers_resumen = ["Total Riesgos", "Abiertos", "Mitigados", "Cerrados", "Críticos", "Altos", "Score Promedio"]
    for c, h in enumerate(headers_resumen, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    valores = [metricas["total"], metricas["abiertos"], metricas["mitigados"],
               metricas["cerrados"], metricas["criticos"], metricas["altos"], metricas["score_promedio"]]
    for c, v in enumerate(valores, 1):
        celda = ws.cell(row, c, v)
        _aplicar_estilo(celda, font=FONT_BODY, border=THIN_BORDER, alignment=ALIGN_CENTER)

    # Detalle de riesgos
    row += 2
    _merge_y_escribir(ws, row, 1, row, 11, "DETALLE DE RIESGOS", FONT_SECCION)
    row += 1
    headers = ["ID", "Descripción", "Probabilidad", "Impacto", "Score", "Clasificación",
               "Estado", "Mitigación", "Responsable", "Fecha Ident.", "Notas"]
    for c, h in enumerate(headers, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1

    # Ordenar por score descendente
    riesgos_ordenados = sorted(riesgos, key=lambda r: r.get("score", 0), reverse=True)

    for r in riesgos_ordenados:
        fill_estado = None
        if r["estado"] == "Abierto":
            fill_estado = FILL_ABIERTO
        elif r["estado"] == "Mitigado":
            fill_estado = FILL_MITIGADO
        elif r["estado"] == "Cerrado":
            fill_estado = FILL_CERRADO

        # Fill de clasificación
        _, fill_clasif = clasificar_riesgo(r.get("score", 0))

        vals = [
            r["id"], r["descripcion"], r["probabilidad"], r["impacto"],
            r.get("score", ""), r.get("clasificacion", ""),
            r["estado"], r["mitigacion"], r["responsable"],
            r["fecha_identificacion"], r["notas"]
        ]
        for c, v in enumerate(vals, 1):
            celda = ws.cell(row, c, v)
            fill = fill_clasif if c == 6 else fill_estado
            _aplicar_estilo(celda, font=FONT_BODY, border=THIN_BORDER, fill=fill,
                           alignment=ALIGN_CENTER if c in (1, 3, 4, 5, 6, 7, 10) else None)
        row += 1


def main():
    print("Leyendo riesgos...")
    riesgos = leer_riesgos()

    if not riesgos:
        print("No se encontraron riesgos. Asegúrese de que la hoja 'Riesgos' exista en el Excel.")
        return

    metricas = calcular_metricas(riesgos)
    fecha_gen = datetime.now().strftime("%d/%m/%Y")

    print(f"  Total: {metricas['total']}")
    print(f"  Abiertos: {metricas['abiertos']}")
    print(f"  Críticos: {metricas['criticos']}")
    print(f"  Score promedio: {metricas['score_promedio']}")

    try:
        wb = openpyxl.load_workbook(ARCHIVO_SALIDA_REPORTE)
    except FileNotFoundError:
        print(f"Archivo de reporte no encontrado: {ARCHIVO_SALIDA_REPORTE}")
        print("Ejecute primero generar_reporte_nevasa.py")
        return

    agregar_hoja_riesgos(wb, riesgos, metricas, fecha_gen)

    archivo_temp = ARCHIVO_SALIDA_REPORTE.replace(".xlsx", "_temp.xlsx")
    wb.save(archivo_temp)
    try:
        os.replace(archivo_temp, ARCHIVO_SALIDA_REPORTE)
    except PermissionError:
        print(f"\n  [!] No se pudo sobrescribir. Guardado como: {archivo_temp}")
        return
    print(f"\nAnálisis de riesgos agregado a: {ARCHIVO_SALIDA_REPORTE}")


if __name__ == "__main__":
    main()
