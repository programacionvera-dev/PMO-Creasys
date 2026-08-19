"""
Generador de Reporte Ejecutivo GPI CB
Lee 'Seguimiento Requerimientos NevasaCB.xlsx' y genera 'Reporte_Estatus_GPI_CB.xlsx'
"""

import sys
import os
from datetime import datetime
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config.rutas import ARCHIVO_FUENTE, ARCHIVO_SALIDA_REPORTE

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ESTADOS_COMPLETADO = {"Producción", "Certificación Nevasa"}
ESTADOS_EN_PROGRESO = {"Sin Formalizar", "Análisis", "Desarrollo", "Certificación Creasys", "Stand By", "Cancelado"}

# Colores base - Nevasa Corporativo
NEVASA_OSCURO = "FF333F48"
NEVASA_GRIS = "FF97999B"
NEVASA_NARANJA = "FFFA4616"
NEVASA_CLARO = "FFE8E8E8"
BLANCO = "FFFFFFFF"
GRIS_TEXTO = "FF333F48"

# Colores condicionales - Estados
COLOR_SIN_FORMALIZAR = "FFCCCC"
COLOR_PENDIENTE = "FFCCCC"
COLOR_ANALISIS = "FFFFCC"
COLOR_EN_DESARROLLO = "FFFFCC"
COLOR_CERT_CREASYS = "FFFFFFFF"
COLOR_QA_NEVASA = "CCE5FF"
COLOR_PRODUCCION = "CCFFCC"
COLOR_STANDBY = "FFF8DC"
COLOR_CANCELADO = "90EE90"

# Colores condicionales - Prioridad
COLOR_MAXIMA = "FFB6C1"
COLOR_ALTA = "FFDAB9"

# Colores indicador (5 etapas)
INDICADOR_CRITICO_BG = "FFFF4444"
INDICADOR_CRITICO_FG = "FFFFFFFF"
INDICADOR_ALERTA_BG = "FFFF8C00"
INDICADOR_ALERTA_FG = "FFFFFFFF"
INDICADOR_SEGUIMIENTO_BG = "FFFFD700"
INDICADOR_SEGUIMIENTO_FG = "FF000000"
INDICADOR_ACEPTABLE_BG = "FF90EE90"
INDICADOR_ACEPTABLE_FG = "FF000000"
INDICADOR_BUENO_BG = "FF228B22"
INDICADOR_BUENO_FG = "FFFFFFFF"

# Fonts
FONT_TITULO = Font(name="Calibri", size=16, bold=True, color=NEVASA_OSCURO)
FONT_FECHA = Font(name="Calibri", size=10, color=NEVASA_GRIS)
FONT_SECCION = Font(name="Calibri", size=11, bold=True)
FONT_HEADER = Font(name="Calibri", size=9, bold=True, color=BLANCO)
FONT_BODY = Font(name="Calibri", size=9)
FONT_HEADER_DETALLE = Font(name="Calibri", size=10, bold=True, color=BLANCO)
FONT_TITULO_AREA = Font(name="Calibri", size=12, bold=True, color=BLANCO)

# Fills
FILL_HEADER = PatternFill(start_color=NEVASA_OSCURO, end_color=NEVASA_OSCURO, fill_type="solid")
FILL_AREA = PatternFill(start_color=NEVASA_OSCURO, end_color=NEVASA_OSCURO, fill_type="solid")

# Fills condicionales - Estados
FILL_SIN_FORMALIZAR = PatternFill(start_color=COLOR_SIN_FORMALIZAR, end_color=COLOR_SIN_FORMALIZAR, fill_type="solid")
FILL_PENDIENTE = PatternFill(start_color=COLOR_PENDIENTE, end_color=COLOR_PENDIENTE, fill_type="solid")
FILL_ANALISIS = PatternFill(start_color=COLOR_ANALISIS, end_color=COLOR_ANALISIS, fill_type="solid")
FILL_EN_DESARROLLO = PatternFill(start_color=COLOR_EN_DESARROLLO, end_color=COLOR_EN_DESARROLLO, fill_type="solid")
FILL_CERT_CREASYS = PatternFill(start_color=COLOR_CERT_CREASYS, end_color=COLOR_CERT_CREASYS, fill_type="solid")
FILL_QA_NEVASA = PatternFill(start_color=COLOR_QA_NEVASA, end_color=COLOR_QA_NEVASA, fill_type="solid")
FILL_PRODUCCION = PatternFill(start_color=COLOR_PRODUCCION, end_color=COLOR_PRODUCCION, fill_type="solid")
FILL_STANDBY = PatternFill(start_color=COLOR_STANDBY, end_color=COLOR_STANDBY, fill_type="solid")
FILL_CANCELADO = PatternFill(start_color=COLOR_CANCELADO, end_color=COLOR_CANCELADO, fill_type="solid")

# Fills condicionales - Prioridad
FILL_MAXIMA = PatternFill(start_color=COLOR_MAXIMA, end_color=COLOR_MAXIMA, fill_type="solid")
FILL_ALTA = PatternFill(start_color=COLOR_ALTA, end_color=COLOR_ALTA, fill_type="solid")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left")

THIN_BORDER = Border(
    left=Side(style="thin", color="FFB0B0B0"),
    right=Side(style="thin", color="FFB0B0B0"),
    top=Side(style="thin", color="FFB0B0B0"),
    bottom=Side(style="thin", color="FFB0B0B0"),
)

# Mapa de colores por estado
ESTADO_FILL_MAP = {
    "Sin Formalizar": FILL_SIN_FORMALIZAR,
    "Pendiente": FILL_PENDIENTE,
    "Análisis": FILL_ANALISIS,
    "Desarrollo": FILL_EN_DESARROLLO,
    "Certificación Creasys": FILL_CERT_CREASYS,
    "Certificación Nevasa": FILL_QA_NEVASA,
    "Producción": FILL_PRODUCCION,
    "Stand By": FILL_STANDBY,
    "Cancelado": FILL_CANCELADO,
}

# Mapa de colores por prioridad
PRIORIDAD_FILL_MAP = {
    "MÁXIMA": FILL_MAXIMA,
    "ALTA": FILL_ALTA,
}

# Orden personalizado para distribución por estado
ORDEN_ESTADOS = [
    "Producción",
    "Certificación Nevasa",
    "Certificación Creasys",
    "Desarrollo",
    "Análisis",
    "Cancelado",
    "Pendiente"
]


def clasificar_estado(estado):
    if estado in ESTADOS_COMPLETADO:
        return "Completado"
    if estado in ESTADOS_EN_PROGRESO:
        return "En Progreso"
    return "Pendiente"


def obtener_indicador(pct_completados):
    """Retorna (texto, fill_bg, font) según la etapa."""
    if pct_completados < 20:
        return ("CRÍTICO", PatternFill(start_color=INDICADOR_CRITICO_BG, end_color=INDICADOR_CRITICO_BG, fill_type="solid"),
                Font(name="Calibri", size=11, bold=True, color=INDICADOR_CRITICO_FG))
    elif pct_completados < 40:
        return ("ALERTA", PatternFill(start_color=INDICADOR_ALERTA_BG, end_color=INDICADOR_ALERTA_BG, fill_type="solid"),
                Font(name="Calibri", size=11, bold=True, color=INDICADOR_ALERTA_FG))
    elif pct_completados < 60:
        return ("REQUIERE ATENCIÓN", PatternFill(start_color=INDICADOR_SEGUIMIENTO_BG, end_color=INDICADOR_SEGUIMIENTO_BG, fill_type="solid"),
                Font(name="Calibri", size=11, bold=True, color=INDICADOR_SEGUIMIENTO_FG))
    elif pct_completados < 80:
        return ("EN DESARROLLO", PatternFill(start_color=INDICADOR_ACEPTABLE_BG, end_color=INDICADOR_ACEPTABLE_BG, fill_type="solid"),
                Font(name="Calibri", size=11, bold=True, color=INDICADOR_ACEPTABLE_FG))
    else:
        return ("SATISFACTORIO", PatternFill(start_color=INDICADOR_BUENO_BG, end_color=INDICADOR_BUENO_BG, fill_type="solid"),
                Font(name="Calibri", size=11, bold=True, color=INDICADOR_BUENO_FG))


def leer_fuente():
    wb = openpyxl.load_workbook(ARCHIVO_FUENTE, read_only=False)
    ws = wb["Seguimiento Nevasa"]
    filas = []
    for r in range(5, ws.max_row + 1):
        valores = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in valores):
            continue
        
        # Mapeo de columnas según nueva estructura
        indice = valores[0] if len(valores) > 0 and valores[0] else None             # Col 1
        requerimiento = valores[1] if len(valores) > 1 and valores[1] else ""  # Col 2
        detalle = valores[2] if len(valores) > 2 and valores[2] else ""        # Col 3
        area = valores[3] if len(valores) > 3 and valores[3] else ""           # Col 4
        solicitante = valores[4] if len(valores) > 4 and valores[4] else ""    # Col 5
        categoria = valores[5] if len(valores) > 5 and valores[5] else ""      # Col 6
        prioridad_raw = valores[6] if len(valores) > 6 else None              # Col 7
        asignado_a = valores[8] if len(valores) > 8 and valores[8] else ""    # Col 9
        etapa = valores[9] if len(valores) > 9 and valores[9] else "Pendiente"  # Col 10
        fecha_recepcion = valores[11] if len(valores) > 11 else None          # Col 12
        real_inicio = valores[17] if len(valores) > 17 else None              # Col 18
        fecha_real_qa = valores[18] if len(valores) > 18 else None            # Col 19
        aprobado_pap = valores[19] if len(valores) > 19 and valores[19] else ""  # Col 20
        fecha_produccion = valores[20] if len(valores) > 20 else None         # Col 21
        historia = valores[25] if len(valores) > 25 and valores[25] else ""   # Col 26
        periodo = valores[27] if len(valores) > 27 else None                  # Col 28
        estimado_fin_desarrollo = valores[14] if len(valores) > 14 else None  # Col 15 (O)
        
        # Mapeo de prioridad
        if prioridad_raw is None or prioridad_raw == "" or prioridad_raw == "Sin Prioridad":
            prioridad = "SP"
        elif prioridad_raw == 0 or prioridad_raw == "0":
            prioridad = "MÁXIMA"
        elif prioridad_raw == 1 or prioridad_raw == "1":
            prioridad = "ALTA"
        else:
            prioridad = str(prioridad_raw).upper()
        
        estado_clasificado = clasificar_estado(etapa)
        
        periodo_str = ""
        if periodo:
            if isinstance(periodo, datetime):
                periodo_str = periodo.strftime("%Y-%m")
            else:
                periodo_str = str(periodo)
        
        fecha_recepcion_str = "Sin definir"
        if fecha_recepcion:
            if isinstance(fecha_recepcion, datetime):
                fecha_recepcion_str = fecha_recepcion.strftime("%d/%m/%Y")
            else:
                fecha_recepcion_str = str(fecha_recepcion)
        
        fecha_real_qa_str = "Sin definir"
        if fecha_real_qa:
            if isinstance(fecha_real_qa, datetime):
                fecha_real_qa_str = fecha_real_qa.strftime("%d/%m/%Y")
            else:
                fecha_real_qa_str = str(fecha_real_qa)
        
        fecha_inicio_str = "Sin definir"
        if real_inicio:
            if isinstance(real_inicio, datetime):
                fecha_inicio_str = real_inicio.strftime("%d/%m/%Y")
            else:
                fecha_inicio_str = str(real_inicio)
        
        fecha_produccion_str = "Sin definir"
        if fecha_produccion:
            if isinstance(fecha_produccion, datetime):
                fecha_produccion_str = fecha_produccion.strftime("%d/%m/%Y")
            else:
                fecha_produccion_str = str(fecha_produccion)
        
        estimado_fin_desarrollo_str = "Sin definir"
        if estimado_fin_desarrollo:
            if isinstance(estimado_fin_desarrollo, datetime):
                estimado_fin_desarrollo_str = estimado_fin_desarrollo.strftime("%d/%m/%Y")
            else:
                estimado_fin_desarrollo_str = str(estimado_fin_desarrollo)
        
        filas.append({
            "indice": indice,
            "aprobado_pap": aprobado_pap,
            "periodo": periodo_str,
            "tipo": str(categoria) if categoria else "Sin tipo",
            "prioridad": str(prioridad).upper(),
            "area": area,
            "fecha_solicitud": fecha_recepcion_str,
            "quien_solicita": solicitante,
            "requerimiento": requerimiento,
            "detalle_requerimiento": detalle,
            "fecha_inicio_ejecucion": fecha_inicio_str,
            "fecha_entrega_qa": fecha_real_qa_str,
            "fecha_entrega_produccion": fecha_produccion_str,
            "estimado_fin_desarrollo": estimado_fin_desarrollo_str,
            "estado_raw": str(etapa),
            "estado_clasificado": estado_clasificado,
            "historia": historia,
        })
    wb.close()
    return filas


def calcular_metricas(filas):
    total = len(filas)
    completados = sum(1 for f in filas if f["estado_clasificado"] == "Completado")
    en_progreso = sum(1 for f in filas if f["estado_clasificado"] == "En Progreso")
    pendientes = sum(1 for f in filas if f["estado_clasificado"] == "Pendiente")
    porcentaje = f"{(completados / total * 100):.0f}%" if total > 0 else "0%"
    pct_numero = round((completados / total * 100)) if total > 0 else 0

    por_estado = defaultdict(int)
    for f in filas:
        por_estado[f["estado_raw"]] += 1

    areas_unicas = []
    for f in filas:
        if f["area"] and f["area"] not in areas_unicas:
            areas_unicas.append(f["area"])

    por_area = {}
    for area in areas_unicas:
        items = [f for f in filas if f["area"] == area]
        por_area[area] = {
            "total": len(items),
            "pendiente": sum(1 for f in items if f["estado_clasificado"] == "Pendiente"),
            "en_progreso": sum(1 for f in items if f["estado_clasificado"] == "En Progreso"),
            "completado": sum(1 for f in items if f["estado_clasificado"] == "Completado"),
        }

    periodos = sorted(set(f["periodo"] for f in filas if f["periodo"]))
    por_periodo = {}
    for p in periodos:
        items = [f for f in filas if f["periodo"] == p]
        por_periodo[p] = {
            "total": len(items),
            "pendiente": sum(1 for f in items if f["estado_clasificado"] == "Pendiente"),
            "en_progreso": sum(1 for f in items if f["estado_clasificado"] == "En Progreso"),
            "completado": sum(1 for f in items if f["estado_clasificado"] == "Completado"),
        }

    por_prioridad = defaultdict(int)
    for f in filas:
        por_prioridad[f["prioridad"]] += 1

    por_tipo = defaultdict(int)
    por_tipo_estado = defaultdict(lambda: defaultdict(int))
    for f in filas:
        por_tipo[f["tipo"]] += 1
        por_tipo_estado[f["tipo"]][f["estado_raw"]] += 1

    return {
        "total": total,
        "completados": completados,
        "en_progreso": en_progreso,
        "pendientes": pendientes,
        "porcentaje": porcentaje,
        "pct_numero": pct_numero,
        "por_estado": dict(por_estado),
        "por_area": por_area,
        "por_periodo": por_periodo,
        "por_prioridad": dict(por_prioridad),
        "por_tipo": dict(por_tipo),
        "por_tipo_estado": {k: dict(v) for k, v in por_tipo_estado.items()},
    }


def _aplicar_estilo_celda(celda, font=None, fill=None, alignment=None, border=None):
    if font:
        celda.font = font
    if fill:
        celda.fill = fill
    if alignment:
        celda.alignment = alignment
    if border:
        celda.border = border


def _merge_y_escribir(ws, r1, c1, r2, c2, valor, font, alignment=None, fill=None):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    celda = ws.cell(r1, c1, valor)
    _aplicar_estilo_celda(celda, font=font, alignment=alignment or ALIGN_LEFT, fill=fill)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if r != r1 or c != c1:
                cell = ws.cell(r, c)
                _aplicar_estilo_celda(cell, font=font, fill=fill)


def _obtener_fill_estado(estado):
    return ESTADO_FILL_MAP.get(estado)


def _obtener_fill_prioridad(prioridad):
    return PRIORIDAD_FILL_MAP.get(prioridad)


def crear_resumen_ejecutivo(wb, metricas, fecha_gen):
    ws = wb.create_sheet("Resumen Ejecutivo")
    ws.column_dimensions["A"].width = 30
    for col in ["B", "C", "D", "E", "F", "G"]:
        ws.column_dimensions[col].width = 13

    # Título
    _merge_y_escribir(ws, 1, 1, 1, 9, "REPORTE DE ESTATUS - SISTEMA GPI CB",
                       FONT_TITULO, ALIGN_CENTER)
    _merge_y_escribir(ws, 2, 1, 2, 9,
                       f"Fecha de Generación: {fecha_gen}",
                       FONT_FECHA, ALIGN_CENTER)

    # 1. Métricas Generales
    row = 4
    _merge_y_escribir(ws, row, 1, row, 9, "1. MÉTRICAS GENERALES", FONT_SECCION)
    row += 1
    headers = ["Total Requerimientos", "", "Completados", "", "En Progreso", "", "Pendientes", "", "% Cumplimiento"]
    for c, h in enumerate(headers, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    valores = [metricas["total"], None, metricas["completados"], None,
               metricas["en_progreso"], None, metricas["pendientes"], None,
               metricas["pct_numero"] / 100]
    for c, v in enumerate(valores, 1):
        celda = ws.cell(row, c, v)
        _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER,
                              alignment=ALIGN_CENTER if v is not None else None)
        if c == 9:  # Formato porcentaje para columna % Cumplimiento
            celda.number_format = '0%'

    # Disclaimer sobre Certificación Nevasa
    row += 2
    _merge_y_escribir(ws, row, 1, row, 9, 
        "NOTA: Los requerimientos en 'Certificación Nevasa' se consideran completados "
        "ya que para Creasys representan entrega final del desarrollo.",
        Font(name="Calibri", size=9, italic=True, color=NEVASA_GRIS))

    # 2. Estado General del Proyecto
    row += 2
    _merge_y_escribir(ws, row, 1, row, 9, "2. ESTADO GENERAL DEL PROYECTO", FONT_SECCION)
    row += 1
    celda_label = ws.cell(row, 1, "Indicador:")
    _aplicar_estilo_celda(celda_label, font=FONT_BODY)
    texto_indicador, fill_indicador, font_indicador = obtener_indicador(metricas["pct_numero"])
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
    celda_valor = ws.cell(row, 3, texto_indicador)
    _aplicar_estilo_celda(celda_valor, font=font_indicador, fill=fill_indicador, alignment=ALIGN_CENTER)
    for c in range(4, 6):
        ws.cell(row, c).fill = fill_indicador

    # Escala de indicadores (columna M)
    escala_row = 8
    FONT_ESCALA_TITLE = Font(name="Calibri", size=10, bold=True)
    FONT_ESCALA_BODY = Font(name="Calibri", size=9)

    _merge_y_escribir(ws, escala_row, 13, escala_row, 15, "ESCALA DE INDICADORES", FONT_ESCALA_TITLE)
    escala_row += 1
    for c, h in enumerate(["Rango", "Indicador"], 13):
        celda = ws.cell(escala_row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    escala_row += 1

    escalas = [
        ("< 20%", "CRÍTICO", INDICADOR_CRITICO_BG, INDICADOR_CRITICO_FG),
        ("20% - 39%", "ALERTA", INDICADOR_ALERTA_BG, INDICADOR_ALERTA_FG),
        ("40% - 59%", "REQUIERE ATENCIÓN", INDICADOR_SEGUIMIENTO_BG, INDICADOR_SEGUIMIENTO_FG),
        ("60% - 79%", "EN DESARROLLO", INDICADOR_ACEPTABLE_BG, INDICADOR_ACEPTABLE_FG),
        (">= 80%", "SATISFACTORIO", INDICADOR_BUENO_BG, INDICADOR_BUENO_FG),
    ]
    for rango, texto, bg, fg in escalas:
        fill_escala = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        font_escala = Font(name="Calibri", size=9, bold=True, color=fg)
        celda_rango = ws.cell(escala_row, 13, rango)
        _aplicar_estilo_celda(celda_rango, font=FONT_ESCALA_BODY, border=THIN_BORDER, alignment=ALIGN_CENTER)
        celda_texto = ws.cell(escala_row, 14, texto)
        _aplicar_estilo_celda(celda_texto, font=font_escala, fill=fill_escala, border=THIN_BORDER, alignment=ALIGN_CENTER)
        escala_row += 1

    ws.column_dimensions["M"].width = 14
    ws.column_dimensions["N"].width = 20

    # 3. Distribución por Estado
    row += 2
    _merge_y_escribir(ws, row, 1, row, 9, "3. DISTRIBUCIÓN POR ESTADO", FONT_SECCION)
    row += 1
    for c, h in enumerate(["Estado", "Cantidad", "% del Total"], 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    estados_conocidos = [e for e in ORDEN_ESTADOS if e in metricas["por_estado"]]
    estados_nuevos = [e for e in metricas["por_estado"] if e not in ORDEN_ESTADOS]
    estados_orden = estados_conocidos + sorted(estados_nuevos)
    for estado in estados_orden:
        cant = metricas["por_estado"][estado]
        pct = cant / metricas['total'] if metricas['total'] > 0 else 0
        fill_estado = _obtener_fill_estado(estado)
        for c, v in enumerate([estado, cant, pct], 1):
            celda = ws.cell(row, c, v)
            _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER, fill=fill_estado,
                                  alignment=ALIGN_CENTER if c > 1 else None)
            if c == 3:  # Formato porcentaje
                celda.number_format = '0%'
        row += 1

    # 4. Distribución por Área
    row += 1
    _merge_y_escribir(ws, row, 1, row, 9, "4. DISTRIBUCIÓN POR ÁREA", FONT_SECCION)
    row += 1
    for c, h in enumerate(["Área", "Total", "Pendiente", "En Progreso", "Completado"], 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    for area, datos in metricas["por_area"].items():
        vals = [area, datos["total"], datos["pendiente"], datos["en_progreso"], datos["completado"]]
        for c, v in enumerate(vals, 1):
            celda = ws.cell(row, c, v)
            _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER,
                                  alignment=ALIGN_CENTER if c > 1 else None)
        row += 1

    # 5. Distribución por Período
    row += 1
    _merge_y_escribir(ws, row, 1, row, 9, "5. DISTRIBUCIÓN POR PERÍODO", FONT_SECCION)
    row += 1
    for c, h in enumerate(["Período", "Total", "Pendiente", "En Progreso", "Completado"], 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    for periodo, datos in metricas["por_periodo"].items():
        vals = [periodo, datos["total"], datos["pendiente"], datos["en_progreso"], datos["completado"]]
        for c, v in enumerate(vals, 1):
            celda = ws.cell(row, c, v)
            _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER,
                                  alignment=ALIGN_CENTER if c > 1 else None)
        row += 1

    # 6. Distribución por Prioridad
    row += 1
    _merge_y_escribir(ws, row, 1, row, 9, "6. DISTRIBUCIÓN POR PRIORIDAD", FONT_SECCION)
    row += 1
    for c, h in enumerate(["Prioridad", "Cantidad", "% del Total"], 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1
    orden_prio_dist = ["MÁXIMA", "ALTA", "MEDIA/ALTA", "MEDIA", "BAJA", "SP"]
    prioridades_existentes = [p for p in orden_prio_dist if p in metricas["por_prioridad"]]
    for prioridad in prioridades_existentes:
        cant = metricas["por_prioridad"][prioridad]
        pct = cant / metricas['total'] if metricas['total'] > 0 else 0
        fill_prio = _obtener_fill_prioridad(prioridad)
        for c, v in enumerate([prioridad, cant, pct], 1):
            celda = ws.cell(row, c, v)
            _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER, fill=fill_prio,
                                  alignment=ALIGN_CENTER if c > 1 else None)
            if c == 3:  # Formato porcentaje
                celda.number_format = '0%'
        row += 1

    # 7. Distribución por Tipo (Matriz Tipo × Estado)
    row += 1
    _merge_y_escribir(ws, row, 1, row, 9, "7. DISTRIBUCIÓN POR TIPO", FONT_SECCION)
    row += 1
    headers_tipo = ["Tipo", "Total", "% Producción", "Pendiente", "Certificación Nevasa", "Desarrollo", "Producción"]
    for c, h in enumerate(headers_tipo, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER, fill=FILL_HEADER, border=THIN_BORDER)
    row += 1

    tipos_orden = ["Incidencia", "Solicitud", "Requerimiento", "Mejora", "Evolutivo", "Proceso Manual"]
    FILL_TOTAL = PatternFill(start_color=NEVASA_CLARO, end_color=NEVASA_CLARO, fill_type="solid")
    FONT_ROJO = Font(name="Calibri", size=9, color="FFFF0000", bold=True)
    FONT_AMARILLO = Font(name="Calibri", size=9, color="FFCC8000", bold=True)
    FONT_VERDE = Font(name="Calibri", size=9, color="FF008000", bold=True)
    FONT_NEGRITA = Font(name="Calibri", size=9, bold=True)

    totales_tipo = {"total": 0, "produccion": 0, "pendiente": 0, "cert_qa": 0, "desarrollo": 0}

    for tipo in tipos_orden:
        if tipo not in metricas["por_tipo"]:
            continue
        cant_total = metricas["por_tipo"][tipo]
        estados = metricas["por_tipo_estado"].get(tipo, {})
        cant_pendiente = estados.get("Pendiente", 0) + estados.get("Sin Formalizar", 0)
        cant_cert_qa = estados.get("Certificación Nevasa", 0)
        cant_desarrollo = estados.get("Desarrollo", 0) + estados.get("Análisis", 0)
        cant_produccion = estados.get("Producción", 0)
        pct_prod = (cant_produccion / cant_total * 100) if cant_total > 0 else 0

        totales_tipo["total"] += cant_total
        totales_tipo["produccion"] += cant_produccion
        totales_tipo["pendiente"] += cant_pendiente
        totales_tipo["cert_qa"] += cant_cert_qa
        totales_tipo["desarrollo"] += cant_desarrollo

        # Color de % Producción según umbral
        if pct_prod < 40:
            font_pct = FONT_ROJO
        elif pct_prod < 80:
            font_pct = FONT_AMARILLO
        else:
            font_pct = FONT_VERDE

        # Negrita para Incidencia y Solicitud
        font_tipo = FONT_NEGRITA if tipo in ("Incidencia", "Solicitud") else FONT_BODY

        vals = [tipo, cant_total, pct_prod / 100, cant_pendiente, cant_cert_qa, cant_desarrollo, cant_produccion]
        for c, v in enumerate(vals, 1):
            celda = ws.cell(row, c, v)
            font_cell = font_pct if c == 3 else font_tipo
            _aplicar_estilo_celda(celda, font=font_cell, border=THIN_BORDER,
                                  alignment=ALIGN_CENTER if c > 1 else None)
            if c == 3:  # Formato porcentaje para columna % Producción
                celda.number_format = '0%'
        row += 1

    # Fila de totales
    pct_total_prod = (totales_tipo["produccion"] / totales_tipo["total"]) if totales_tipo["total"] > 0 else 0
    vals_total = ["TOTAL", totales_tipo["total"], pct_total_prod,
                  totales_tipo["pendiente"], totales_tipo["cert_qa"], totales_tipo["desarrollo"], totales_tipo["produccion"]]
    for c, v in enumerate(vals_total, 1):
        celda = ws.cell(row, c, v)
        _aplicar_estilo_celda(celda, font=Font(name="Calibri", size=9, bold=True), border=THIN_BORDER,
                              fill=FILL_TOTAL, alignment=ALIGN_CENTER if c > 1 else None)
        if c == 3:  # Formato porcentaje para columna % Producción
            celda.number_format = '0%'
    row += 1


def crear_detalle_por_area(wb, filas, fecha_gen):
    ws = wb.create_sheet("Detalle por área")
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 18
    ws.column_dimensions["K"].width = 18
    ws.column_dimensions["L"].width = 15
    ws.column_dimensions["M"].width = 18
    ws.column_dimensions["N"].width = 12
    ws.column_dimensions["O"].width = 45

    _merge_y_escribir(ws, 1, 1, 1, 15, "DETALLE DE REQUERIMIENTOS POR ÁREA",
                       FONT_TITULO, ALIGN_CENTER)
    _merge_y_escribir(ws, 2, 1, 2, 15, f"Fecha: {fecha_gen}",
                       FONT_FECHA, ALIGN_CENTER)

    items_no_produccion = [f for f in filas if f["estado_raw"] != "Producción"]

    def _normalizar_area(area):
        if not area:
            return "Sin área"
        if "Comercial" in area:
            return "Área Comercial y Clientes Web"
        elif "Operaciones" in area:
            return "Área Operaciones"
        elif "normativos" in area:
            return "Temas normativos"
        return area

    orden_prioridad = {"MÁXIMA": 0, "ALTA": 1, "MEDIA/ALTA": 2, "MEDIA": 3, "BAJA": 4, "SP": 5}
    orden_estado = {
        "Sin Formalizar": 0,
        "Pendiente": 1,
        "Análisis": 2,
        "Desarrollo": 3,
        "Certificación Creasys": 4,
        "Certificación Nevasa": 5,
        "Producción": 6,
        "Stand By": 7,
        "Cancelado": 8
    }
    orden_area = {
        "Área Comercial y Clientes Web": 0,
        "Área Operaciones": 1,
        "Temas normativos": 2
    }

    for item in items_no_produccion:
        item["_area_norm"] = _normalizar_area(item["area"])

    items_no_produccion.sort(key=lambda f: (
        orden_area.get(f["_area_norm"], 9),
        orden_estado.get(f["estado_raw"], 5),
        orden_prioridad.get(f["prioridad"], 5)
    ))

    row = 4
    headers = ["#", "Área", "Prioridad", "Requerimiento", "Detalle Requerimiento", "Solicitante",
               "Estado", "Fecha Solicitud", "Fecha Inicio Ejecución", "Estimado Fin Desarrollo",
               "Fecha Entrega QA", "Aprobado PaP", "Fecha Entrega Producción", "Periodo", "Observaciones"]

    for c, h in enumerate(headers, 1):
        celda = ws.cell(row, c, h)
        _aplicar_estilo_celda(celda, font=FONT_HEADER_DETALLE, fill=FILL_HEADER,
                              border=THIN_BORDER)
    row += 1

    for item in items_no_produccion:
        vals = [
            item["indice"],
            item["_area_norm"],
            item["prioridad"],
            item["requerimiento"],
            item["detalle_requerimiento"],
            item["quien_solicita"],
            item["estado_raw"],
            item["fecha_solicitud"],
            item["fecha_inicio_ejecucion"],
            item["estimado_fin_desarrollo"],
            item["fecha_entrega_qa"],
            item["aprobado_pap"] if item["estado_raw"] in ("Certificación Nevasa", "Producción") else "",
            item["fecha_entrega_produccion"],
            item["periodo"],
            item["historia"],
        ]
        fill_estado = _obtener_fill_estado(item["estado_raw"])
        for c, v in enumerate(vals, 1):
            celda = ws.cell(row, c, v)
            if c in (1, 3, 7, 8, 9, 10, 11, 12, 13, 14):
                alineacion = ALIGN_CENTER
            elif c in (4, 5, 15):
                alineacion = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif c == 6:
                alineacion = Alignment(horizontal="center", vertical="center")
            else:
                alineacion = None
            _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER,
                                  fill=fill_estado, alignment=alineacion)
        fill_prio = _obtener_fill_prioridad(item["prioridad"])
        if fill_prio:
            ws.cell(row, 3).fill = fill_prio
        row += 1

    ws.auto_filter.ref = f"A4:O{row - 1}"


def crear_requerimientos_produccion(wb, filas, fecha_gen):
    ws = wb.create_sheet("Requerimientos en Producción")
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 12
    ws.column_dimensions["I"].width = 45

    _merge_y_escribir(ws, 1, 1, 1, 9, "REQUERIMIENTOS EN PRODUCCIÓN",
                       FONT_TITULO, ALIGN_CENTER)
    _merge_y_escribir(ws, 2, 1, 2, 9, f"Fecha: {fecha_gen}",
                       FONT_FECHA, ALIGN_CENTER)

    filas_produccion = [f for f in filas if f["estado_raw"] == "Producción"]

    por_area = defaultdict(list)
    for f in filas_produccion:
        area = f["area"] if f["area"] else "Sin área"
        if "Comercial" in area:
            area_normalizada = "Area Comercial y Clientes Web"
        elif "Operaciones" in area:
            area_normalizada = "Area Operaciones"
        elif "normativos" in area:
            area_normalizada = "Temas normativos"
        else:
            area_normalizada = area
        por_area[area_normalizada].append(f)

    orden_prioridad = {"MÁXIMA": 0, "ALTA": 1, "MEDIA/ALTA": 2, "MEDIA": 3, "BAJA": 4, "SP": 5}

    row = 4
    headers = ["#", "Prioridad", "Requerimiento", "Detalle Requerimiento", "Solicitante",
               "Estado", "Fecha Producción", "Periodo", "Observaciones"]
    nombre_areas = {
        "Area Comercial y Clientes Web": "AREA COMERCIAL Y CLIENTES WEB",
        "\u00c1rea Comercial y Clientes Web": "AREA COMERCIAL Y CLIENTES WEB",
        "Area Operaciones": "AREA OPERACIONES",
        "\u00c1rea Operaciones": "AREA OPERACIONES",
        "Temas normativos": "TEMAS NORMATIVOS",
    }

    for area_key in ["Area Comercial y Clientes Web", "Area Operaciones", "Temas normativos"]:
        if area_key not in por_area:
            continue
        items = por_area[area_key]
        items.sort(key=lambda f: orden_prioridad.get(f["prioridad"], 5))

        titulo = f"--- {nombre_areas.get(area_key, area_key.upper())} ---"
        _merge_y_escribir(ws, row, 1, row, 9, titulo, FONT_TITULO_AREA, ALIGN_LEFT, fill=FILL_AREA)
        row += 1

        for c, h in enumerate(headers, 1):
            celda = ws.cell(row, c, h)
            _aplicar_estilo_celda(celda, font=FONT_HEADER_DETALLE, fill=FILL_HEADER,
                                  border=THIN_BORDER)
        row += 1

        for idx, item in enumerate(items, 1):
            vals = [
                idx,
                item["prioridad"],
                item["requerimiento"],
                item["detalle_requerimiento"],
                item["quien_solicita"],
                item["estado_raw"],
                item["fecha_entrega_produccion"],
                item["periodo"],
                item["historia"],
            ]
            fill_estado = _obtener_fill_estado(item["estado_raw"])
            for c, v in enumerate(vals, 1):
                celda = ws.cell(row, c, v)
                if c in (1, 2, 6, 7, 8):
                    alineacion = ALIGN_CENTER
                elif c in (3, 4, 9):
                    alineacion = Alignment(horizontal="left", vertical="center", wrap_text=True)
                elif c == 5:
                    alineacion = Alignment(horizontal="center", vertical="center")
                else:
                    alineacion = None
                _aplicar_estilo_celda(celda, font=FONT_BODY, border=THIN_BORDER,
                                      fill=fill_estado, alignment=alineacion)
            fill_prio = _obtener_fill_prioridad(item["prioridad"])
            if fill_prio:
                ws.cell(row, 2).fill = fill_prio
            row += 1

        row += 1


def main():
    print(f"Leyendo fuente: {ARCHIVO_FUENTE}")
    filas = leer_fuente()
    print(f"  Filas leídas: {len(filas)}")

    metricas = calcular_metricas(filas)
    fecha_gen = datetime.now().strftime("%d/%m/%Y")

    print(f"  Total: {metricas['total']}")
    print(f"  Completados: {metricas['completados']}")
    print(f"  En Progreso: {metricas['en_progreso']}")
    print(f"  Pendientes: {metricas['pendientes']}")
    print(f"  % Cumplimiento: {metricas['porcentaje']}")
    texto_ind, _, _ = obtener_indicador(metricas["pct_numero"])
    print(f"  Indicador: {texto_ind}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    crear_resumen_ejecutivo(wb, metricas, fecha_gen)
    crear_detalle_por_area(wb, filas, fecha_gen)
    crear_requerimientos_produccion(wb, filas, fecha_gen)

    archivo_temp = ARCHIVO_SALIDA_REPORTE.replace(".xlsx", "_temp.xlsx")
    wb.save(archivo_temp)
    try:
        os.replace(archivo_temp, ARCHIVO_SALIDA_REPORTE)
    except PermissionError:
        print(f"\n  [!] No se pudo sobrescribir (archivo abierto?). Guardado como: {archivo_temp}")
        return
    print(f"\nReporte generado: {ARCHIVO_SALIDA_REPORTE}")


if __name__ == "__main__":
    main()
