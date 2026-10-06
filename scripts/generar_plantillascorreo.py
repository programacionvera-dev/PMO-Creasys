import openpyxl
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config.rutas import ARCHIVO_FUENTE, CARPETA_SALIDA_TEMPLATES, URL_FIRMA, MAPEO_COLUMNAS

def mapear_prioridad(valor):
    """Mapea el valor del Excel a texto de prioridad"""
    if valor == 0 or valor == '0':
        return 'MÁXIMA'
    elif valor == 1 or valor == '1':
        return 'ALTA'
    elif valor is None or str(valor).strip() == '' or str(valor).strip() == 'Sin Prioridad':
        return 'Sin Prioridad'
    else:
        return str(valor)

MESES_ESPANOL = {
    1: 'ENERO', 2: 'FEBRERO', 3: 'MARZO', 4: 'ABRIL',
    5: 'MAYO', 6: 'JUNIO', 7: 'JULIO', 8: 'AGOSTO',
    9: 'SEPTIEMBRE', 10: 'OCTUBRE', 11: 'NOVIEMBRE', 12: 'DICIEMBRE'
}

def construir_placeholder_plan_prueba(nombre_archivo):
    """Construye el placeholder para Power Automate"""
    if not nombre_archivo:
        return None
    
    return f"{{{{PLAN:{nombre_archivo}}}}}"

def leer_datos():
    """Lee el Excel y retorna ítems por estado"""
    wb = openpyxl.load_workbook(ARCHIVO_FUENTE, read_only=True)
    ws = wb['Seguimiento Nevasa']
    
    # Usar mapeo de columnas desde configuración
    m = MAPEO_COLUMNAS
    
    qa_items = []
    prod_items = []
    todos_items = []
    
    for r in range(5, ws.max_row + 1):
        etapa = ws.cell(r, m["etapa"] + 1).value  # +1 porque openpyxl usa base 1
        if not etapa:
            continue
            
        item = {
            'indice': ws.cell(r, m['indice'] + 1).value,
            'periodo': ws.cell(r, m["periodo"] + 1).value,
            'categoria': ws.cell(r, m["categoria"] + 1).value,
            'prioridad': ws.cell(r, m["prioridad"] + 1).value,
            'area': ws.cell(r, m["area"] + 1).value,
            'fecha_recepcion': ws.cell(r, m["fecha_recepcion"] + 1).value,
            'solicitante': ws.cell(r, m["solicitante"] + 1).value,
            'requerimiento': ws.cell(r, m["requerimiento"] + 1).value,
            'detalle': ws.cell(r, m["detalle"] + 1).value,
            'asignado_a': ws.cell(r, m["asignado"] + 1).value,
            'fecha_qa': ws.cell(r, m["fecha_real_qa"] + 1).value,
            'aprobado_pap': ws.cell(r, m['aprobado_pap'] + 1).value if 'aprobado_pap' in m else None,
            'plan_prueba': ws.cell(r, m["plan_prueba"] + 1).value,
            'etapa': etapa,
            'historia': ws.cell(r, m["historia"] + 1).value,
            'fecha_produccion': ws.cell(r, m["fecha_produccion"] + 1).value,
        }
        
        todos_items.append(item)
        
        # QA: considerar solo desarrollos con aprobado_pap vacío y etapa Certificación Nevasa o Certificación Creasys
        etapa_str = str(etapa).strip() if etapa else ''
        aprobado_pap_vacio = not item.get('aprobado_pap') or str(item.get('aprobado_pap')).strip() == ''
        if etapa_str == 'Certificación Nevasa' and aprobado_pap_vacio:
            qa_items.append(item)
        elif etapa_str == 'Certificación Creasys' and aprobado_pap_vacio:
            qa_items.append(item)
        # Produccion: solo incluir si NO tiene fecha Produccion (aun no estan en Produccion)
        elif etapa_str == 'Producción' and not item['fecha_produccion']:
            prod_items.append(item)
    
    wb.close()
    return qa_items, prod_items, todos_items

def generar_html_qa(items, fecha):
    """Genera plantilla HTML para QA"""
    rows_html = ""
    for i, item in enumerate(items, 1):
        prioridad = mapear_prioridad(item['prioridad'])
        placeholder = construir_placeholder_plan_prueba(item['plan_prueba'])
        
        if placeholder:
            plan_prueba_html = f'<a href="{placeholder}" style="color:#2F5496;text-decoration:underline" target="_blank">Descargar Plan de Prueba</a>'
        else:
            plan_prueba_html = '-'
        
        rows_html += f"""
        <tr>
            <td style="border:1px solid #ddd;padding:8px;text-align:center">{item.get('indice') or ''}</td>
            <td style="border:1px solid #ddd;padding:8px;text-align:center">{prioridad}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['area'] or ''}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['requerimiento'] or ''}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['detalle'] or ''}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['solicitante'] or ''}</td>
            <td style="border:1px solid #ddd;padding:8px;text-align:center">{plan_prueba_html}</td>
        </tr>"""
    
    return f"""<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"></head>
<body style="font-family:Calibri,sans-serif;margin:20px;color:#333f48">
    <h2 style="color:#333f48;font-size:20px">&#x1F4CB; Requerimientos Disponibles Certificación Nevasa</h2>
    <p>Estimados:</p>
    <p>Le informamos que los siguientes requerimientos se encuentran disponibles para su validación en el ambiente de Certificación Nevasa:</p>
    
    <p><strong>Fecha:</strong> {fecha}</p>
    <p><strong>Total de requerimientos:</strong> {len(items)}</p>
    
    <table style="border-collapse:collapse;width:100%;margin:20px 0">
        <thead>
            <tr style="background:#333f48;color:white">
                <th style="border:1px solid #ddd;padding:8px">Indice</th>
                <th style="border:1px solid #ddd;padding:8px">Prioridad</th>
                <th style="border:1px solid #ddd;padding:8px">Área</th>
                <th style="border:1px solid #ddd;padding:8px">Requerimiento</th>
                <th style="border:1px solid #ddd;padding:8px">Detalle</th>
                <th style="border:1px solid #ddd;padding:8px">Solicitante</th>
                <th style="border:1px solid #ddd;padding:8px">Plan de Prueba</th>
            </tr>
        </thead>
        <tbody>{rows_html}</tbody>
    </table>
    
    <div style="background:#f5f5f5;padding:15px;border-left:4px solid #fa4616;margin:20px 0">
        <p style="margin:0 0 10px 0"><strong style="color:#fa4616">Importante:</strong> Les solicitamos su pronta coordinación para realizar las pruebas de los requerimientos indicados para que el despliegue a producción se realice sin complicaciones y en el menor tiempo posible.</p>
        <p style="margin:0 0 10px 0">Si un requerimento aun se encuentra en Certificación Nevasa es debido a que aun no se certifica o contiene actualizaciones. Favor revisar el detalle del requerimiento en el "Reporte_Estatus_GPI_CB" según su Indice.</p>
        <p style="margin:0 0 10px 0">Para los requerimientos que contengan Plan de Prueba, estos pueden descargarse para su uso como guía para las pruebas a realizar.</p>
    </div>
    
    <p>Quedamos atentos a sus comentarios.</p>
    
    <div style="margin-top:30px;border-top:1px solid #97999B;padding-top:10px">
        <img src="{URL_FIRMA}" alt="Creasys" style="height:250px">
    </div>
</body>
</html>"""

def generar_html_produccion(items, fecha):
    """Genera plantilla HTML para Produccion"""
    rows_html = ""
    for i, item in enumerate(items, 1):
        rows_html += f"""
        <tr>
            <td style="border:1px solid #ddd;padding:8px;text-align:center">{item.get('indice') or ''}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['area'] or ''}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['requerimiento'] or ''}</td>
            <td style="border:1px solid #ddd;padding:8px">{item['solicitante'] or ''}</td>
        </tr>"""
    
    return f"""<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"></head>
<body style="font-family:Calibri,sans-serif;margin:20px;color:#333f48">
    <h2 style="color:#333f48;font-size:20px">&#x1F680; Requerimientos a Producción</h2>
    <p>Estimado equipo:</p>
    <p>Le informamos que los siguientes requerimientos serán desplegados en el ambiente de producción:</p>
    
    <p><strong>Fecha:</strong> {fecha}</p>
    <p><strong>Total de requerimientos:</strong> {len(items)}</p>
    
    <table style="border-collapse:collapse;width:100%;margin:20px 0">
        <thead>
            <tr style="background:#333f48;color:white">
                <th style="border:1px solid #ddd;padding:8px">Indice</th>
                <th style="border:1px solid #ddd;padding:8px">Área</th>
                <th style="border:1px solid #ddd;padding:8px">Requerimiento</th>
                <th style="border:1px solid #ddd;padding:8px">Solicitante</th>
            </tr>
        </thead>
        <tbody>{rows_html}</tbody>
    </table>
    
    <div style="background:#f5f5f5;padding:15px;border-left:4px solid #fa4616;margin:20px 0">
        <h3 style="margin-top:0;color:#333f48">Detalles del despliegue</h3>
        <p style="margin:5px 0"><strong>Ventana de integración y validación:</strong> 15 a 30 minutos</p>
        <p style="margin:10px 0 0 0;padding:10px;background:#fff3cd;border-left:4px solid #ffc107">
            <strong>Recomendación:</strong> Durante el despliegue a Producción, se sugiere que los sistemas de GPI 
            involucrados no se encuentren en uso por parte de los usuarios, ya que esto podría generar 
            complicaciones en la integración.
        </p>
    </div>
    
    <div style="margin-top:30px;border-top:1px solid #97999B;padding-top:10px">
        <img src="{URL_FIRMA}" alt="Creasys" style="height:250px">
    </div>
</body>
</html>"""

def generar_html_reporte_estatus(metricas, fecha):
    """Genera plantilla HTML para Reporte de Estatus"""
    return f"""<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"></head>
<body style="font-family:Calibri,sans-serif;margin:20px;color:#333f48">
    <h2 style="color:#333f48;font-size:20px">&#x1F4CA; Reporte de Estatus Actualizado</h2>
    <p>Estimados:</p>
    <p>Le informamos que el Reporte de Estatus de Seguimiento Nevasa ha sido actualizado con la información más reciente a la fecha de envío.</p>
    
    <p><strong>Fecha:</strong> {fecha}</p>
    
    <div style="background:#f5f5f5;padding:15px;border-left:4px solid #2F5496;margin:20px 0">
        <h3 style="margin-top:0;color:#333f48">Resumen del Estado Actual</h3>
        <table style="border-collapse:collapse;width:100%;margin:10px 0">
            <tr>
                <td style="padding:5px 15px 5px 0"><strong>Total de requerimientos:</strong></td>
                <td style="padding:5px 0">{metricas['total']}</td>
            </tr>
            <tr>
                <td style="padding:5px 15px 5px 0"><strong>Completados:</strong></td>
                <td style="padding:5px 0;color:#00B050">{metricas['completados']}</td>
            </tr>
            <tr>
                <td style="padding:5px 15px 5px 0"><strong>En Progreso:</strong></td>
                <td style="padding:5px 0;color:#ED7D31">{metricas['en_progreso']}</td>
            </tr>
            <tr>
                <td style="padding:5px 15px 5px 0"><strong>Pendientes:</strong></td>
                <td style="padding:5px 0;color:#A5A5A5">{metricas['pendientes']}</td>
            </tr>
            <tr>
                <td style="padding:5px 15px 5px 0"><strong>% Cumplimiento:</strong></td>
                <td style="padding:5px 0;font-weight:bold">{metricas['porcentaje']}</td>
            </tr>
        </table>
    </div>
    
    <div style="background:#f5f5f5;padding:15px;border-left:4px solid #fa4616;margin:20px 0">
        <p style="margin:0"><strong style="color:#fa4616">Importante:</strong> Le solicitamos revisar el reporte adjunto para mantenerse al tanto del avance de los requerimientos asignados a su área.</p>
    </div>
    
    <p>Quedamos atentos a sus comentarios.</p>
    
</body>
</html>"""


def main():
    os.makedirs(CARPETA_SALIDA_TEMPLATES, exist_ok=True)
    
    qa_items, prod_items, todos_items = leer_datos()
    fecha = datetime.now().strftime('%d/%m/%Y')
    
    # Calcular métricas (misma lógica que generar_reporte_nevasa.py)
    ESTADOS_COMPLETADO = {"Producción", "Certificación Nevasa"}
    ESTADOS_EN_PROGRESO = {"Sin Formalizar", "Análisis", "Desarrollo", "Certificación Creasys", "Stand By", "Cancelado"}
    
    total = len(todos_items)
    completados = sum(1 for item in todos_items if item['etapa'] in ESTADOS_COMPLETADO)
    en_progreso = sum(1 for item in todos_items if item['etapa'] in ESTADOS_EN_PROGRESO)
    pendientes = sum(1 for item in todos_items if item['etapa'] not in ESTADOS_COMPLETADO and item['etapa'] not in ESTADOS_EN_PROGRESO)
    porcentaje = f"{(completados / total * 100):.0f}%" if total > 0 else "0%"
    
    metricas = {
        'total': total,
        'completados': completados,
        'en_progreso': en_progreso,
        'pendientes': pendientes,
        'porcentaje': porcentaje
    }
    
    # Generar HTML QA
    html_qa = generar_html_qa(qa_items, fecha)
    ruta_qa = os.path.join(CARPETA_SALIDA_TEMPLATES, 'qa.html')
    with open(ruta_qa, 'w', encoding='utf-8') as f:
        f.write(html_qa)
    
    # Generar HTML Produccion
    html_prod = generar_html_produccion(prod_items, fecha)
    ruta_prod = os.path.join(CARPETA_SALIDA_TEMPLATES, 'produccion.html')
    with open(ruta_prod, 'w', encoding='utf-8') as f:
        f.write(html_prod)
    
    # Generar HTML Reporte Estatus
    html_estatus = generar_html_reporte_estatus(metricas, fecha)
    ruta_estatus = os.path.join(CARPETA_SALIDA_TEMPLATES, 'reporte_estatus.html')
    with open(ruta_estatus, 'w', encoding='utf-8') as f:
        f.write(html_estatus)
    
    print(f"Plantillas generadas en: {CARPETA_SALIDA_TEMPLATES}")
    print(f"  - qa.html ({len(qa_items)} items)")
    print(f"  - produccion.html ({len(prod_items)} items)")
    print(f"  - reporte_estatus.html (total: {total}, completados: {completados}, en_progreso: {en_progreso}, pendientes: {pendientes})")
    
    return ruta_qa, ruta_prod, ruta_estatus

if __name__ == '__main__':
    main()
